import csv
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from apps.core.export import Column, Table, csv_response, xlsx_response


def table() -> Table:
    return Table(
        "Test",
        [
            Column("Datum", "date"),
            Column("Text"),
            Column("Částka", "money"),
            Column("Podíl", "percent"),
            Column("Počet", "int"),
        ],
        [[date(2026, 10, 5), "Jídlo; a „uvozovky“", 123_450, 12.34, 3], [None, "", 5, None, 0]],
    )


def test_csv_for_czech_excel() -> None:
    content = csv_response([table()], "test").content
    assert content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == ["Datum", "Text", "Částka", "Podíl", "Počet"]
    assert rows[1] == ["5.10.2026", "Jídlo; a „uvozovky“", "1234,50", "12,3", "3"]
    assert rows[2] == ["", "", "0,05", "", "0"]


def test_xlsx_typed_cells() -> None:
    response = xlsx_response(
        [table(), Table("Druhý list s dlouhým názvem přes 31 znaků", [Column("A")], [["x"]])],
        "test",
    )
    assert response["Content-Disposition"] == 'attachment; filename="test.xlsx"'
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Test", "Druhý list s dlouhým názvem pře"]
    sheet = book["Test"]
    assert [c.value for c in sheet[1]] == ["Datum", "Text", "Částka", "Podíl", "Počet"]
    assert sheet[1][0].font.bold
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:E3"
    row = sheet[2]
    assert row[0].value == datetime(2026, 10, 5)
    assert Decimal(str(row[2].value)) == Decimal("1234.5") and "Kč" in row[2].number_format
    assert row[4].value == 3

"""Export tabulek do CSV (český Excel) a XLSX (P14).

Částky přicházejí jako int v haléřích a do souborů jdou v Kč – v CSV jako text
s desetinnou čárkou, v XLSX jako Decimal s formátem Kč (bez plovoucí čárky).
"""

import csv
import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any, Literal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

Kind = Literal["text", "date", "money", "int", "percent"]
FORMATS = ("csv", "xlsx")
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NUMBER_FORMATS = {
    "date": "d.m.yyyy",
    "money": '#,##0.00 "Kč"',
    "int": "0",
    "percent": "0.0 %",
}
WIDTHS = {"text": 24, "date": 12, "money": 16, "int": 10, "percent": 10}


@dataclass(frozen=True)
class Column:
    title: str
    kind: Kind = "text"
    width: int | None = None


@dataclass
class Table:
    name: str  # název listu v XLSX (max. 31 znaků)
    columns: Sequence[Column]
    rows: Iterable[Sequence[Any]] = field(default_factory=list)


def money(cents: int) -> Decimal:
    return Decimal(cents) / 100


def _csv_value(value: Any, kind: Kind) -> str:
    if value is None or value == "":
        return ""
    if kind == "money":
        return f"{money(value):.2f}".replace(".", ",")
    if kind == "date":
        return f"{value.day}.{value.month}.{value.year}"
    if kind == "percent":
        return f"{value:.1f}".replace(".", ",")
    return str(value)


def _xlsx_value(value: Any, kind: Kind) -> Any:
    if value is None or value == "":
        return None
    if kind == "money":
        return money(value)
    if kind == "percent":
        return Decimal(str(round(value, 3))) / 100
    return value


def _download(content: bytes, filename: str, content_type: str) -> HttpResponse:
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def csv_response(tables: Sequence[Table], filename: str) -> HttpResponse:
    """CSV pro český Excel: UTF-8 s BOM, středník, desetinná čárka.
    Víc tabulek se oddělí prázdným řádkem a řádkem s názvem tabulky."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    for index, table in enumerate(tables):
        if len(tables) > 1:
            if index:
                writer.writerow([])
            writer.writerow([table.name])
        writer.writerow([c.title for c in table.columns])
        for row in table.rows:
            writer.writerow(
                [_csv_value(v, c.kind) for v, c in zip(row, table.columns, strict=True)]
            )
    return _download(
        ("﻿" + buffer.getvalue()).encode("utf-8"), f"{filename}.csv", "text/csv; charset=utf-8"
    )


def _header(sheet: WriteOnlyWorksheet, title: str) -> Cell:
    cell = WriteOnlyCell(sheet, value=title)
    cell.font = Font(bold=True)
    return cell


def xlsx_response(tables: Sequence[Table], filename: str) -> HttpResponse:
    """XLSX: list na tabulku, tučná hlavička, zamrzlý první řádek, automatický filtr."""
    workbook = Workbook(write_only=True)
    for table in tables:
        sheet = workbook.create_sheet(title=table.name[:31])
        sheet.freeze_panes = "A2"
        for index, column in enumerate(table.columns, start=1):
            width = column.width or WIDTHS[column.kind]
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.append([_header(sheet, c.title) for c in table.columns])
        count = 0
        for row in table.rows:
            cells = []
            for value, column in zip(row, table.columns, strict=True):
                cell = WriteOnlyCell(sheet, value=_xlsx_value(value, column.kind))
                if column.kind in NUMBER_FORMATS:
                    cell.number_format = NUMBER_FORMATS[column.kind]
                cells.append(cell)
            sheet.append(cells)
            count += 1
        last = get_column_letter(len(table.columns))
        sheet.auto_filter.ref = f"A1:{last}{count + 1}"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return _download(buffer.getvalue(), f"{filename}.xlsx", XLSX_TYPE)


def export_response(fmt: str, tables: Sequence[Table], filename: str) -> HttpResponse:
    if fmt == "xlsx":
        return xlsx_response(tables, filename)
    return csv_response(tables, filename)

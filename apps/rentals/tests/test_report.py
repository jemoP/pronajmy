"""Výkaz (P12), podklad pro daně (P13, ZM-004), export výkazu (P14) a upozornění (P11)."""

import csv
import io
from datetime import date

import pytest
from django.test import Client
from openpyxl import load_workbook

from apps.core.formatting import format_czk
from apps.money.models import Account, Category, CategoryCode, CategoryKind, EntryKind
from apps.money.services.entries import EntryInput, create_entry, result_for
from apps.rentals.models import (
    Communication,
    Inspection,
    InspectionPlan,
    InspectionType,
    Lease,
    Property,
    Repair,
)
from apps.rentals.services.alerts import OVER, WARNING, rental_alerts
from apps.rentals.services.payments import record_payment
from apps.rentals.services.report import flat_rate, rental_report, years_overview

TODAY = date(2026, 10, 16)


def add(
    kind: EntryKind,
    amount: int,
    account: Account,
    day: date,
    code: CategoryCode | None = None,
    prop: Property | None = None,
) -> None:
    category = Category.objects.get(code=code).id if code else None
    create_entry(
        EntryInput(
            date=day,
            kind=kind,
            amount_cents=amount,
            account_id=account.pk,
            category_id=category,
            property_id=prop.pk if prop else None,
        )
    )


def fill(lease: Lease, account: Account) -> None:
    """Platby 18 000 Kč za leden–září, jistota, oprava, služby, vrácený přeplatek,
    ostatní příjem bez bytu, výběr a vklad, plánovaný příjem a jeden příjem v roce 2025."""
    flat = lease.property
    for month in range(1, 10):
        day = date(2026, month, 5)
        record_payment(lease, day=day, amount=1_800_000, account_id=account.pk, today=day)
    day = date(2026, 1, 2)
    record_payment(
        lease, day=day, amount=0, deposit_amount=3_600_000, account_id=account.pk, today=day
    )
    add(EntryKind.EXPENSE, 250_000, account, date(2026, 3, 1), CategoryCode.REPAIRS, flat)
    add(EntryKind.EXPENSE, 1_000_000, account, date(2026, 4, 1), CategoryCode.UTILITIES, flat)
    add(EntryKind.EXPENSE, 50_000, account, date(2026, 5, 31), CategoryCode.SETTLEMENT, flat)
    add(EntryKind.INCOME, 100_000, account, date(2026, 6, 1), CategoryCode.OTHER_INCOME)
    add(EntryKind.WITHDRAWAL, 5_000_000, account, date(2026, 6, 30))
    add(EntryKind.DEPOSIT, 1_000_000, account, date(2026, 7, 1))
    add(EntryKind.INCOME, 1_800_000, account, date(2026, 12, 5), CategoryCode.RENT, flat)
    add(EntryKind.INCOME, 1_000_000, account, date(2025, 6, 5), CategoryCode.RENT, flat)


def test_report_matches_entries(lease: Lease, rental_account: Account) -> None:
    """Akceptace etapy 4: výkaz sedí s pohyby (jistota a výběry se nepočítají)."""
    fill(lease, rental_account)
    report = rental_report(2026, TODAY)
    year = result_for(date(2026, 1, 1), date(2026, 12, 31), TODAY)
    assert (report.income, report.expense) == (year.income_cents, year.expense_cents)
    assert report.income == 16_200_000 + 100_000 - 50_000  # vrácený přeplatek snižuje příjmy
    assert report.expense == 250_000 + 1_000_000
    flat, no_property = report.properties
    assert flat.name == "Byt Laštůvkova"
    # Nájem a služby odděleně (ZM-008).
    assert flat.rent.income == {"Nájemné": 13_800_000}
    assert flat.rent.expense == {"Opravy a údržba": 250_000}
    assert flat.services.income == {"Vyúčtování služeb": -50_000, "Zálohy na služby": 2_400_000}
    assert flat.services.expense == {"Služby": 1_000_000}
    assert (no_property.name, no_property.rent.income) == ("Bez bytu", {"Ostatní příjmy": 100_000})
    totals = report.totals
    assert (totals.rent_income, totals.rent_expense, totals.rent_profit) == (
        13_900_000,
        250_000,
        13_650_000,
    )
    assert (totals.service_income, totals.service_expense, totals.service_difference) == (
        2_350_000,
        1_000_000,
        1_350_000,
    )
    assert totals.profit == report.profit == totals.rent_profit + totals.service_difference
    assert report.months[4].service_income == 300_000 - 50_000  # zálohy − vrácený přeplatek
    assert report.months_income[0] == 1_800_000 and report.months_expense[2] == 250_000
    assert report.months_income[4] == 1_800_000 - 50_000
    assert report.months_income[11] == 0  # plánovaný příjem ne
    assert (report.flows.withdrawn, report.flows.deposited) == (5_000_000, 1_000_000)
    # příjmy 16 200 000 + jistota 3 600 000 + ostatní 100 000 + 2025: 1 000 000
    # − náklady 1 300 000 − výběr 5 000 000 + vklad 1 000 000 (plánovaný ne)
    assert report.flows.balance_end == 15_600_000
    assert report.previous.income == 1_000_000
    assert [(y.year, y.income) for y in years_overview(TODAY)] == [
        (2026, report.income),
        (2025, 1_000_000),
    ]


def test_report_for_one_property(lease: Lease, rental_account: Account) -> None:
    fill(lease, rental_account)
    report = rental_report(2026, TODAY, lease.property_id)
    assert [p.name for p in report.properties] == ["Byt Laštůvkova"]
    assert report.income == 16_200_000 - 50_000


def test_tax_summary_excludes_services(lease: Lease, rental_account: Account) -> None:
    """ZM-004: zálohy, vyúčtování a náklad Služby se do daní nepočítají."""
    fill(lease, rental_account)
    tax = rental_report(2026, TODAY).tax
    assert tax.income == 13_800_000 + 100_000
    assert tax.actual_expenses == 250_000
    assert tax.excluded_income == {"Zálohy na služby": 2_400_000, "Vyúčtování služeb": -50_000}
    assert tax.excluded_expense == {"Služby": 1_000_000}
    assert tax.flat_rate == 4_170_000 and tax.better == "paušál"
    assert tax.flat_base == 13_900_000 - 4_170_000
    # Vlastní účel s vypnutým příznakem se z daní vyřadí; výkaz se nemění.
    before = rental_report(2026, TODAY).income
    Category.objects.filter(code=CategoryCode.OTHER_INCOME).update(taxable=False)
    report = rental_report(2026, TODAY)
    assert report.tax.income == 13_800_000 and report.income == before


def test_flat_rate_cap() -> None:
    """Akceptace etapy 4: paušál respektuje strop 600 000 Kč."""
    assert flat_rate(1_000_000_00) == 300_000_00
    assert flat_rate(2_000_000_00) == 600_000_00
    assert flat_rate(3_000_000_00) == 600_000_00
    assert flat_rate(-100) == 0


def test_report_page(auth_client: Client, lease: Lease, rental_account: Account) -> None:
    fill(lease, rental_account)
    html = auth_client.get("/report?year=2026").content.decode()
    assert format_czk(16_250_000) in html and "Podklad pro daňové přiznání" in html
    assert "Nepočítá se do daní" in html and "Zálohy na služby" in html
    assert "Po letech" in html and 'data-chart="rental-chart"' in html
    assert "jan@example.cz" not in html and "777 123 456" not in html
    assert "Výkaz" in auth_client.get("/").content.decode()  # v navigaci
    # Neznámý rok a byt se ignorují.
    assert auth_client.get("/report?year=1900&property=999").status_code == 200


def test_report_export(auth_client: Client, lease: Lease, rental_account: Account) -> None:
    """Akceptace etapy 4: export neobsahuje kontakty (ani čísla účtů)."""
    rental_account.account_number = "2501234567"
    rental_account.save()
    fill(lease, rental_account)
    response = auth_client.get("/report/export.csv?year=2026")
    assert "pronajmy-2026.csv" in response["Content-Disposition"]
    content = response.content.decode("utf-8-sig")
    assert "Byt Laštůvkova" in content and "Zálohy na služby" in content
    for secret in ("Jan Nájemník", "777 123 456", "jan@example.cz", "2501234567"):
        assert secret not in content, secret
    rows = list(csv.reader(io.StringIO(content), delimiter=";"))
    celkem = ["139000,00", "2500,00", "136500,00", "23500,00", "10000,00", "13500,00", "150000,00"]
    assert ["Celkem", *celkem] in rows
    book = load_workbook(io.BytesIO(auth_client.get("/report/export.xlsx?year=2026").content))
    assert book.sheetnames == ["Byty 2026", "Podle účelů", "Měsíce", "Roky", "Podklad pro DP"]
    assert auth_client.get("/report/export.pdf").status_code == 404


def test_category_taxable_flag_in_ui(auth_client: Client, db: None) -> None:
    services = Category.objects.get(code=CategoryCode.SERVICES)
    assert not services.taxable and Category.objects.get(code=CategoryCode.RENT).taxable
    html = auth_client.get("/categories").content.decode()
    assert "Do daní" in html
    auth_client.post(
        f"/categories/{services.pk}/edit",
        {
            "name": services.name,
            "kind": CategoryKind.INCOME,
            "sort": "20",
            "active": "on",
            "taxable": "on",
        },
    )
    services.refresh_from_db()
    assert services.taxable
    auth_client.post("/categories", {"name": "Parkování", "kind": "INCOME", "sort": "1"})
    assert not Category.objects.get(name="Parkování").taxable  # nezaškrtnuto


@pytest.mark.parametrize("url", ["/report", "/report/export.csv"])
def test_report_requires_login(client: Client, db: None, url: str) -> None:
    assert client.get(url)["Location"].startswith("/login/")


def make_alerts(lease: Lease, rental_account: Account) -> None:
    for month in range(1, 10):
        day = date(2026, month, 5)
        record_payment(lease, day=day, amount=1_800_000, account_id=rental_account.pk, today=day)
    plan = InspectionPlan.objects.create(
        property=lease.property, type=InspectionType.objects.get(name="Servis plynového kotle")
    )
    Inspection.objects.create(
        plan=plan, done_date=date(2025, 11, 1), next_due_date=date(2026, 11, 1), result="OK"
    )
    Repair.objects.create(property=lease.property, reported_date=TODAY, title="Zvonek")
    Communication.objects.create(
        property=lease.property,
        occurred_at="2026-10-01T08:00:00Z",
        channel="PHONE",
        direction="IN",
        subject="Klíče",
        follow_up_date=date(2026, 10, 10),
    )
    lease.end_date = date(2026, 12, 1)
    lease.save()


def test_alerts(lease: Lease, rental_account: Account) -> None:
    """Dluh po splatnosti, revize kotle „za 16 dní“ (akceptace etapy 3), konec smlouvy,
    otevřené opravy a komunikace po termínu; nejdřív položky po termínu."""
    make_alerts(lease, rental_account)
    alerts = rental_alerts(TODAY)
    texts = [(a.level, a.text) for a in alerts]
    assert (
        OVER,
        f"Nájem po splatnosti: Byt Laštůvkova – Jan Nájemník dluží {format_czk(2_200_000)}"
        " (od 5. 9. 2026).",
    ) in texts
    assert any(
        level == WARNING and "Servis plynového kotle" in t and "za 16 dní" in t
        for level, t in texts
    )
    assert (WARNING, "Otevřené opravy: 1.") in texts
    assert any(level == OVER and "Vyřídit: Klíče" in t for level, t in texts)
    assert any("končí 1. 12. 2026 (za 46 dní)" in t for _, t in texts)
    levels = [a.level for a in alerts]
    assert levels == sorted(levels, key=lambda level: level != OVER)
    # Uvnitř skupiny podle data: dluh od 5. 9. před komunikací do 10. 10.
    over = [a.text for a in alerts if a.level == OVER]
    assert over[0].startswith("Nájem po splatnosti") and over[1].startswith("Vyřídit")


def test_lease_past_end_not_ended(lease: Lease) -> None:
    lease.end_date = date(2026, 9, 30)
    lease.save()
    texts = [a.text for a in rental_alerts(TODAY)]
    assert any("skončila 30. 9. 2026 – ukonči ji" in t for t in texts)


def test_alert_banner_on_other_pages(
    auth_client: Client, lease: Lease, rental_account: Account
) -> None:
    make_alerts(lease, rental_account)
    for number in range(6):
        Communication.objects.create(
            property=lease.property,
            occurred_at="2026-10-01T08:00:00Z",
            channel="PHONE",
            direction="IN",
            subject=f"Úkol {number}",
            follow_up_date=date(2026, 10, 1),
        )
    html = auth_client.get("/accounts").content.decode()
    assert "Nájem po splatnosti" in html and "Další upozornění:" in html
    assert html.count('class="notice-error"') + html.count('class="notice-warning"') == 5
    overview = auth_client.get("/").content.decode()
    assert "Další upozornění:" not in overview and "Úkol 5" in overview  # celý seznam


def test_no_alerts_without_login(client: Client, lease: Lease) -> None:
    html = client.get("/login/").content.decode()
    assert "Nájem po splatnosti" not in html

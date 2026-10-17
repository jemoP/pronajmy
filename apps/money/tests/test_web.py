import csv
import io

import pytest
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from apps.money.models import Account, Category, CategoryCode, Entry, EntryKind
from apps.rentals.models import Property

from .conftest import add, category

URLS = ("/entries", "/entries/export.csv", "/accounts", "/accounts/new", "/categories")


@pytest.mark.parametrize("url", URLS)
def test_pages_require_login(client: Client, db: None, url: str) -> None:
    response = client.get(url)
    assert response.status_code == 302 and response["Location"].startswith("/login/")


@pytest.mark.parametrize("url", ("/", *URLS[:1], *URLS[2:]))
def test_pages_render(auth_client: Client, bank: Account, flat: Property, url: str) -> None:
    assert auth_client.get(url).status_code == 200


def test_create_entry(auth_client: Client, bank: Account, flat: Property) -> None:
    today = timezone.localdate().isoformat()
    response = auth_client.post(
        "/entries?kind=INCOME",
        {
            "date": today,
            "kind": "INCOME",
            "amount": "18 000",
            "account_id": bank.id,
            "category_id": category(CategoryCode.RENT).id,
            "property_id": flat.id,
            "note": "Nájem",
        },
    )
    entry = Entry.objects.get()
    assert entry.amount_cents == 1_800_000 and entry.property == flat
    assert response.status_code == 302
    assert f"saved={entry.id}" in response["Location"] and "kind=INCOME" in response["Location"]
    page = auth_client.get(response["Location"]).content.decode()
    assert "Pohyb uložen." in page and f'value="{bank.id}" selected' in page


def test_create_entry_errors(auth_client: Client, bank: Account) -> None:
    page = auth_client.post(
        "/entries",
        {"date": "2026-10-16", "kind": "INCOME", "amount": "abc", "account_id": bank.id},
    ).content.decode()
    assert "Pohyb nebyl uložen" in page and "1 234,50" in page
    page = auth_client.post(
        "/entries",
        {"date": "2026-10-16", "kind": "INCOME", "amount": "100", "account_id": bank.id},
    ).content.decode()
    assert "Vyber účel." in page
    assert not Entry.objects.exists()


def test_edit_and_delete_entry(auth_client: Client, bank: Account, cash: Account) -> None:
    entry = add(EntryKind.WITHDRAWAL, 100, bank)
    url = f"/entries/{entry.id}/edit?back=/entries%3Fkind%3DWITHDRAWAL"
    assert "Úprava pohybu" in auth_client.get(url).content.decode()
    response = auth_client.post(
        url,
        {
            "date": "2026-10-01",
            "kind": "TRANSFER",
            "amount": "50",
            "account_id": bank.id,
            "to_account_id": cash.id,
        },
    )
    assert response["Location"] == f"/entries?kind=WITHDRAWAL&saved={entry.id}"
    entry.refresh_from_db()
    assert entry.kind == EntryKind.TRANSFER and entry.to_account == cash
    assert auth_client.get(f"/entries/{entry.id}/delete").status_code == 405
    response = auth_client.post(f"/entries/{entry.id}/delete?back=https://evil.example/")
    assert response["Location"] == "/entries?deleted=1"
    assert not Entry.objects.exists()


def test_export_csv_without_account_numbers(
    auth_client: Client, bank: Account, flat: Property
) -> None:
    add(EntryKind.INCOME, 1_800_050, bank, code=CategoryCode.RENT, prop=flat, note="Září; nájem")
    add(EntryKind.WITHDRAWAL, 500_000, bank)
    response = auth_client.get("/entries/export.csv?kind=INCOME")
    content = response.content
    assert content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == [
        "Datum",
        "Druh",
        "Účel",
        "Byt",
        "Částka",
        "Účet",
        "Na účet",
        "Poznámka",
        "Plánováno",
    ]
    assert len(rows) == 2
    assert rows[1][1:8] == [
        "Příjem",
        "Nájemné",
        "Byt Brno 2+kk",
        "18000,50",
        "Fio pronájmy",
        "",
        "Září; nájem",
    ]
    assert bank.account_number not in content.decode("utf-8-sig")


def test_export_xlsx(auth_client: Client, bank: Account) -> None:
    add(EntryKind.DEPOSIT, 100, bank)
    response = auth_client.get("/entries/export.xlsx")
    sheet = load_workbook(io.BytesIO(response.content))["Pohyby"]
    assert sheet.max_row == 2 and sheet["B2"].value == "Vklad z domácnosti"
    assert auth_client.get("/entries/export.pdf").status_code == 404


def test_account_number_masked(auth_client: Client, bank: Account) -> None:
    for url in ("/accounts", f"/accounts/{bank.id}/edit", "/"):
        page = auth_client.get(url).content.decode()
        assert bank.account_number not in page, url
    assert "****4567/2010" in auth_client.get("/accounts").content.decode()


def test_account_form(auth_client: Client, bank: Account) -> None:
    data = {
        "name": "Fio",
        "kind": "BANK",
        "bank_code": "2010",
        "opening": "1 000,50",
        "account_number": "",
        "active": "on",
    }
    assert auth_client.post(f"/accounts/{bank.id}/edit", data).status_code == 302
    bank.refresh_from_db()
    assert bank.opening_balance_cents == 100_050 and bank.account_number == "2501234567"
    page = auth_client.post("/accounts/new", {**data, "opening": "x"}).content.decode()
    assert "1 234,50" in page and Account.objects.count() == 1


def test_categories_page(auth_client: Client, db: None) -> None:
    response = auth_client.post(
        "/categories", {"name": "Parkování", "kind": "INCOME", "sort": "40"}
    )
    assert response.status_code == 302
    own = Category.objects.get(name="Parkování")
    assert own.active and own.sort == 40
    page = auth_client.post("/categories", {"name": "parkování", "kind": "INCOME"}).content
    assert "už existuje" in page.decode()
    rent = category(CategoryCode.RENT)
    auth_client.post(f"/categories/{rent.id}/delete")
    assert Category.objects.filter(pk=rent.pk).exists()
    auth_client.post(f"/categories/{own.id}/delete")
    assert not Category.objects.filter(pk=own.pk).exists()
    response = auth_client.post(
        f"/categories/{rent.id}/edit", {"name": "Nájem", "kind": "INCOME", "sort": "1"}
    )
    rent.refresh_from_db()
    assert response.status_code == 302 and rent.name == "Nájem" and not rent.active


def test_list_sums_match_result_and_show_deposits(
    auth_client: Client, bank: Account, flat: Property
) -> None:
    """Příjmy − Náklady = Výsledek; jistota je v součtech zvlášť (ne v příjmech a nákladech)."""
    day = timezone.localdate()
    add(EntryKind.INCOME, 1_800_000, bank, day=day, code=CategoryCode.RENT, prop=flat)
    add(EntryKind.INCOME, 3_600_000, bank, day=day, code=CategoryCode.DEPOSIT, prop=flat)
    add(EntryKind.EXPENSE, 1_000_000, bank, day=day, code=CategoryCode.DEPOSIT, prop=flat)
    add(EntryKind.EXPENSE, 300_000, bank, day=day, code=CategoryCode.REPAIRS, prop=flat)
    html = auth_client.get("/entries").content.decode()
    nbsp = " "
    assert f"18{nbsp}000,00{nbsp}Kč" in html  # příjmy bez jistoty
    assert f"3{nbsp}000,00{nbsp}Kč" in html  # náklady bez vrácené jistoty
    assert f"15{nbsp}000,00{nbsp}Kč" in html  # výsledek
    assert f"přijato 36{nbsp}000,00{nbsp}Kč · vráceno 10{nbsp}000,00{nbsp}Kč" in html
    assert f"54{nbsp}000,00{nbsp}Kč" not in html  # dřívější součet všech příjmů
